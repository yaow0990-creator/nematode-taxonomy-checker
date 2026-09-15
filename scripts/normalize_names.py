#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Step 1 of the pipeline: turn a name-only table into resolvable Latin genus names.

Input
-----
A CSV/XLSX whose columns may be some subset of:
    中文名 | 拉丁名 | 科 | 目 | 纲
Only 中文名 and/or 拉丁名 are required. Rows that also carry 科/目/纲 are
passed through untouched so step 2 can check them.

What it does
------------
1.  Latin name given      -> normalise, verify against the Nemaplex genus index.
                             If absent, propose near matches (difflib) and mark
                             the row as a suspected spelling error.
2.  Chinese name given    -> resolve via data/cn_latin_genus.json (curated).
3.  Unresolved Chinese    -> written to out/needs_cn_mapping.csv. These must be
                             filled in by a human (or proposed by the model and
                             then confirmed). Nothing is invented here.
4.  Both given            -> cross-checked; disagreement is flagged rather than
                             silently overwritten.

Every row gets a `confidence` label. Anything below `dictionary` is a
proposal, not a fact, and is surfaced in the output as 待人工确认.

Usage
-----
    python scripts/normalize_names.py --input names.csv --out out/normalized.json
    python scripts/normalize_names.py --input table.xlsx --sheet Sheet1
"""

from __future__ import annotations

import argparse
import csv
import difflib
import json
import os
import re
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, "data")

CN_KEYS = ["中文名", "中文属名", "中文", "chinese", "cn", "name_cn"]
LAT_KEYS = ["拉丁名", "拉丁属名", "属名", "拉丁", "latin", "genus", "name_latin"]
FAM_KEYS = ["科", "科名", "科拉丁", "family"]
ORD_KEYS = ["目", "目名", "order"]
CLS_KEYS = ["纲", "纲名", "class"]

CONF_RANK = {
    "exact": 0,
    "dictionary": 1,
    "cn-conflict": 2,
    "spelling-suspect": 3,
    "unresolved": 4,
}


def load_json(name: str):
    with open(os.path.join(DATA, name), encoding="utf-8") as fh:
        return json.load(fh)


def norm_latin(s: str) -> str:
    s = (s or "").strip().strip(".").replace("　", " ")
    s = re.sub(r"\s+", " ", s)
    return s[:1].upper() + s[1:] if s else s


def norm_cn(s: str) -> str:
    s = re.sub(r"\s+", "", s or "")
    return s[:-1] if s.endswith("属") else s


def pick(row: dict, keys: list[str]) -> str:
    lowered = {re.sub(r"\s+", "", str(k)).lower(): v for k, v in row.items() if k is not None}
    for k in keys:
        v = lowered.get(re.sub(r"\s+", "", k).lower())
        if v is not None and str(v).strip():
            return str(v).strip()
    return ""


def read_rows(path: str, sheet: str | None = None) -> list[dict]:
    if path.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook  # optional dependency
        wb = load_workbook(path, data_only=True)
        ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
        header = [c.value for c in ws[1]]
        out = []
        for r in ws.iter_rows(min_row=2, values_only=True):
            if all(v in (None, "") for v in r):
                continue
            out.append({header[i]: r[i] for i in range(min(len(header), len(r)))})
        return out
    with open(path, encoding="utf-8-sig", newline="") as fh:
        return list(csv.DictReader(fh))


def resolve(rows: list[dict], fuzzy_cutoff: float = 0.86) -> tuple[list[dict], list[dict]]:
    genera = load_json("genus_index.json")
    cn_dict = load_json("cn_latin_genus.json")["entries"]
    lookup = {k.lower(): k for k in genera}
    index_keys = list(lookup.keys())

    out: list[dict] = []
    unresolved: list[dict] = []

    for i, row in enumerate(rows, start=1):
        cn_in = pick(row, CN_KEYS)
        lat_in = pick(row, LAT_KEYS)
        rec = {
            "no": i,
            "cn_in": cn_in,
            "latin_in": lat_in,
            "latin": "",
            "latin_nemaplex": "",
            "confidence": "unresolved",
            "note": "",
            "fam_in": pick(row, FAM_KEYS),
            "ord_in": pick(row, ORD_KEYS),
            "cls_in": pick(row, CLS_KEYS),
        }

        # --- 1. Latin name supplied: verify directly -------------------------
        if lat_in:
            cand = norm_latin(lat_in)
            hit = lookup.get(cand.lower())
            if hit:
                rec["latin"] = rec["latin_nemaplex"] = hit
                rec["confidence"] = "exact"
            else:
                close = difflib.get_close_matches(cand.lower(), index_keys, n=3, cutoff=fuzzy_cutoff)
                if close:
                    best = lookup[close[0]]
                    rec["latin"] = best
                    rec["latin_nemaplex"] = best
                    rec["confidence"] = "spelling-suspect"
                    rec["note"] = (f"拉丁名 {cand} 未收录，疑似拼写错误；"
                                   f"最接近 {best}"
                                   + (f"，其他候选：{', '.join(lookup[c] for c in close[1:])}" if len(close) > 1 else ""))
                else:
                    rec["latin"] = cand
                    rec["confidence"] = "unresolved"
                    rec["note"] = f"拉丁名 {cand} 未收录于 Nemaplex，且无相近候选"

        # --- 2./3. Chinese name: resolve through the curated dictionary ------
        if cn_in:
            entry = cn_dict.get(norm_cn(cn_in))
            if entry:
                dict_latin = entry["latin"]
                if rec["latin"] and rec["latin"].lower() != dict_latin.lower():
                    rec["confidence"] = "cn-conflict"
                    rec["note"] = (rec["note"] + "；" if rec["note"] else "") + \
                        f"中文名对应字典值为 {dict_latin}，与表内拉丁名不一致，请人工裁决"
                else:
                    rec["latin"] = dict_latin
                    if not rec["latin_nemaplex"]:
                        rec["latin_nemaplex"] = lookup.get(dict_latin.lower(), dict_latin)
                    if rec["confidence"] in ("unresolved",):
                        rec["confidence"] = "dictionary"
                    if entry.get("confidence") == "seed-unresolved":
                        rec["note"] = (rec["note"] + "；" if rec["note"] else "") + \
                            "该属在 Nemaplex 中无对应记录（种子标为未解决）"
            else:
                if not rec["latin"]:
                    rec["note"] = f"中文名「{cn_in}」不在字典中，需人工补拉丁名"
                else:
                    rec["note"] = (rec["note"] + "；" if rec["note"] else "") + \
                        f"中文名「{cn_in}」不在字典中，已按表内拉丁名处理"
                # Anything that did not reach `exact` still needs the Chinese
                # name mapped by hand - the Latin name is not trustworthy either.
                if rec["confidence"] != "exact":
                    unresolved.append({
                        "no": i, "cn": cn_in,
                        "note": "字典未收录" + ("，且拉丁名未在 Nemaplex 命中" if rec["latin"] else ""),
                    })

        if not rec["latin"] and not cn_in and not lat_in:
            rec["note"] = "空行或名称列缺失"
        out.append(rec)

    return out, unresolved


def main() -> int:
    ap = argparse.ArgumentParser(description="Normalise nematode genus names (step 1).")
    ap.add_argument("--input", required=True)
    ap.add_argument("--sheet", default=None)
    ap.add_argument("--out", default=os.path.join(ROOT, "out", "normalized.json"))
    ap.add_argument("--cutoff", type=float, default=0.86)
    args = ap.parse_args()

    if not os.path.exists(os.path.join(DATA, "genus_index.json")):
        print("data/genus_index.json is missing - run scripts/build_snapshot.py first", file=sys.stderr)
        return 1

    rows = read_rows(args.input, args.sheet)
    resolved, unresolved = resolve(rows, args.cutoff)

    os.makedirs(os.path.dirname(args.out), exist_ok=True)
    with open(args.out, "w", encoding="utf-8") as fh:
        json.dump(resolved, fh, ensure_ascii=False, indent=1)

    pending = os.path.join(os.path.dirname(args.out), "needs_cn_mapping.csv")
    if unresolved:
        with open(pending, "w", encoding="utf-8-sig", newline="") as fh:
            w = csv.DictWriter(fh, fieldnames=["no", "cn", "note", "latin_proposed", "confirmed_by"])
            w.writeheader()
            for u in unresolved:
                w.writerow({**u, "latin_proposed": "", "confirmed_by": ""})

    counts: dict[str, int] = {}
    for r in resolved:
        counts[r["confidence"]] = counts.get(r["confidence"], 0) + 1
    print(f"rows: {len(resolved)} -> {args.out}")
    for k in sorted(counts, key=lambda x: CONF_RANK.get(x, 99)):
        print(f"  {k}: {counts[k]}")
    if unresolved:
        print(f"待人工补充中文名映射: {len(unresolved)} -> {pending}")
    flagged = [r for r in resolved if r["confidence"] in ("spelling-suspect", "cn-conflict", "unresolved")]
    if flagged:
        print("需人工确认的行:")
        for r in flagged[:20]:
            print(f"  [{r['no']}] {r['cn_in']} / {r['latin_in']} -> {r['note']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

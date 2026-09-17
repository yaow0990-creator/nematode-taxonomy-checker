#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Step 4 of the pipeline: render the matched rows into deliverables.

Produces
--------
out/线虫分类核查结果.xlsx   4 sheets
   核查结果      全字段表（含中文译名、功能团、判定、备注）
   待人工确认    需要人过一眼的行
   功能团汇总    functional guild 分布 + c-p 分布 + 取食类群分布
   说明          字段释义、置信度分级、体系说明、数据出处
out/report.html             单文件 HTML 报告

Usage
-----
    python scripts/export.py --input out/matched.json
    python scripts/export.py --input out/matched.json --xlsx out/table.xlsx --no-html
"""

from __future__ import annotations

import argparse
import collections
import datetime as _dt
import html as _html
import importlib.util
import json
import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, "data")

COLUMNS = [
    ("no",                     "序号"),
    ("cn_in",                  "中文名（源表）"),
    ("latin_in",               "拉丁名（源表）"),
    ("latin",                  "采用拉丁名"),
    ("latin_nemaplex",         "Nemaplex 拉丁名"),
    ("fam_in",                 "科（源表）"),
    ("fam_nem",                "科（Nemaplex）"),
    ("fam_alt",                "科（Nemaplex 备选）"),
    ("subcls_nem",             "亚纲"),
    ("subcls_cn",              "亚纲（中文）"),
    ("ord_in",                 "目（源表）"),
    ("ord_nem",                "目（Nemaplex）"),
    ("ord_cn",                 "目（中文）"),
    ("cls_in",                 "纲（源表）"),
    ("cls_nem",                "纲（现代体系）"),
    ("cls_cn",                 "纲（中文）"),
    ("cls_trad",               "纲（经典体系）"),
    ("cls_trad_cn",            "纲（经典·中文）"),
    ("cp",                     "c-p 值"),
    ("cp_cn",                  "c-p 类型"),
    ("feeding_group",          "取食类群码"),
    ("feeding_group_cn",       "取食类群"),
    ("functional_guild",       "功能团"),
    ("functional_guild_cn",    "功能团（中文）"),
    ("putative_feeding",       "站点取食习性"),
    ("putative_feeding_cn",    "站点取食习性（中文）"),
    ("internal_code",          "属内码"),
    ("genus_url",              "属页链接"),
    ("verdict_cn",             "判定"),
    ("confidence",             "置信度"),
    ("note",                   "备注"),
]

CONF_CN = {
    "exact": "精确（Nemaplex 直接命中）",
    "dictionary": "字典命中",
    "cn-conflict": "中文名与拉丁名冲突，待确认",
    "spelling-suspect": "疑似拼写错误，待确认",
    "unresolved": "未解析，待确认",
}

NEEDS_HUMAN = {"cn-conflict", "spelling-suspect", "unresolved", "not-found"}

WARN_FILL = "FFF3CD"
BAD_FILL = "F8D7DA"


def load_json(path: str):
    with open(path, encoding="utf-8") as fh:
        return json.load(fh)


# --------------------------------------------------------------------------- #
# xlsx
# --------------------------------------------------------------------------- #

def write_xlsx(rows: list[dict], labels: dict, meta: dict, path: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    head_fill = PatternFill("solid", fgColor="2F6F4E")
    head_font = Font(color="FFFFFF", bold=True, size=10)
    warn_fill = PatternFill("solid", fgColor=WARN_FILL)
    bad_fill = PatternFill("solid", fgColor=BAD_FILL)
    thin = Side(style="thin", color="D0D0D0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    wb = Workbook()

    # ---- sheet 1: main table ------------------------------------------------
    ws = wb.active
    ws.title = "核查结果"
    ws.append([cn for _, cn in COLUMNS])
    for c in range(1, len(COLUMNS) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill, cell.font = head_fill, head_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws.freeze_panes = "C2"

    for r in rows:
        ws.append([r.get(k, "") for k, _ in COLUMNS])
        excel_row = ws.max_row
        fill = None
        if r.get("confidence") in NEEDS_HUMAN:
            fill = warn_fill
        if r.get("confidence") == "cn-conflict":
            fill = bad_fill
        for c in range(1, len(COLUMNS) + 1):
            cell = ws.cell(row=excel_row, column=c)
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=(c == len(COLUMNS)))
            if fill:
                cell.fill = fill
        link_cell = ws.cell(row=excel_row, column=[k for k, _ in COLUMNS].index("genus_url") + 1)
        if r.get("genus_url"):
            link_cell.hyperlink = r["genus_url"]
            link_cell.font = Font(color="0563C1", underline="single", size=10)

    widths = {"no": 5, "cn_in": 14, "latin_in": 18, "latin": 18, "latin_nemaplex": 18,
              "fam_in": 14, "fam_nem": 16, "subcls_nem": 13, "subcls_cn": 12,
              "ord_in": 12, "ord_nem": 14, "ord_cn": 10, "cls_in": 12, "cls_nem": 13,
              "cls_cn": 10, "cls_trad": 12, "cls_trad_cn": 12, "cp": 6, "cp_cn": 10,
              "feeding_group": 9, "feeding_group_cn": 16, "functional_guild": 10,
              "functional_guild_cn": 16, "putative_feeding": 18, "putative_feeding_cn": 14,
              "internal_code": 8, "genus_url": 22, "verdict_cn": 20, "confidence": 22, "note": 46}
    for i, (key, _) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(key, 14)
    ws.row_dimensions[1].height = 32
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{ws.max_row}"

    # ---- sheet 2: needs human review ---------------------------------------
    ws2 = wb.create_sheet("待人工确认")
    ws2.append(["序号", "中文名（源表）", "拉丁名（源表）", "采用拉丁名", "置信度", "备注", "人工裁决"])
    for c in range(1, 8):
        cell = ws2.cell(row=1, column=c)
        cell.fill, cell.font = head_fill, head_font
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    flagged = [r for r in rows if r.get("confidence") in NEEDS_HUMAN or r.get("verdict") in ("not-found",)]
    for r in flagged:
        ws2.append([r.get("no"), r.get("cn_in"), r.get("latin_in"), r.get("latin"),
                    CONF_CN.get(r.get("confidence"), r.get("confidence")), r.get("note"), ""])
    for col, w in zip("ABCDEFG", (5, 14, 18, 18, 26, 60, 16)):
        ws2.column_dimensions[col].width = w
    if not flagged:
        ws2.append(["", "", "", "", "", "无待人工确认的行", ""])

    # ---- sheet 3: guild summary --------------------------------------------
    ws3 = wb.create_sheet("功能团汇总")
    ws3.append(["功能团分布（取值食类群首字母 + c-p 值）"])
    ws3.cell(row=1, column=1).font = Font(bold=True)
    ws3.append(["功能团", "中文", "属数"])
    guilds = collections.Counter(r["functional_guild"] for r in rows if r.get("functional_guild"))
    gw = {r["functional_guild"]: r for r in rows if r.get("functional_guild")}
    for g, n in sorted(guilds.items()):
        ws3.append([g, gw[g].get("functional_guild_cn", ""), n])
    ws3.append([])
    ws3.append(["c-p 分布"])
    ws3.cell(row=ws3.max_row, column=1).font = Font(bold=True)
    ws3.append(["c-p 值", "类型", "属数"])
    cps = collections.Counter(r["cp"] for r in rows if r.get("cp"))
    cp_ref = {r["cp"]: r for r in rows if r.get("cp")}
    for cp, n in sorted(cps.items()):
        ws3.append([cp, cp_ref[cp].get("cp_cn", ""), n])
    ws3.append([])
    ws3.append(["取食类群分布"])
    ws3.cell(row=ws3.max_row, column=1).font = Font(bold=True)
    ws3.append(["码", "取食类群", "属数"])
    fgs = collections.Counter(r["feeding_group"] for r in rows if r.get("feeding_group"))
    fg_ref = {r["feeding_group"]: r for r in rows if r.get("feeding_group")}
    for fg, n in sorted(fgs.items()):
        ws3.append([fg, fg_ref[fg].get("feeding_group_cn", ""), n])
    for col, w in zip("ABC", (14, 30, 10)):
        ws3.column_dimensions[col].width = w

    # ---- sheet 4: legend ---------------------------------------------------
    ws4 = wb.create_sheet("说明")
    lines = [
        ("数据出处", ""),
        ("分类与功能团数据", "Nemaplex, UC Davis — http://nemaplex.ucdavis.edu/"),
        ("快照生成时间", meta.get("built_at", "")),
        ("属索引规模", f"{meta.get('genus_count', '')} 属，其中 {meta.get('genus_with_cp', '')} 个有 c-p 值、"
                       f"{meta.get('genus_with_guild', '')} 个有功能团"),
        ("科页规模", f"{meta.get('family_count', '')} 科，其中 {meta.get('family_with_order', '')} 科解析出目"),
        ("中文译名", "非 Nemaplex 内容，来自 data/zh_cn_labels.json，为文献通用译法，投稿前请按您引用的体系复核"),
        ("", ""),
        ("体系说明", ""),
        ("现代体系", "De Ley & Blaxter（SSU rDNA）：纲 = Chromadorea 色矛纲 / Enoplea 刺嘴纲；"
                     "Tylenchida 垫刃目已并入 Rhabditida 小杆目"),
        ("经典体系", "Chitwood (1958)：纲 = Adenophorea 无尾感器纲 / Secernentea 尾感器纲"),
        ("经典纲如何得出", "现代 Rhabditida 吸收了经典体系里 Secernentea 的垫刃目、滑刃目、双胃目、"
                           "圆线目、蛔目、旋尾目等，故「目=Rhabditida → Secernentea」，其余归 Adenophorea"),
        ("", ""),
        ("置信度分级", ""),
    ]
    for k, v in CONF_CN.items():
        lines.append((k, v))
    lines += [
        ("", ""),
        ("c-p 值释义", ""),
    ]
    cp_map = labels.get("cp_classes", {})
    for k in sorted(k for k in cp_map if k.isdigit()):
        lines.append((f"{k} = {cp_map[k]['cn']}", cp_map[k]["desc"]))
    lines += [("", ""), ("取食类群释义", "")]
    for k, v in sorted(labels.get("feeding_groups", {}).items()):
        if k.isdigit():
            lines.append((f"{k}", f"{v['cn']} / {v['en']}"))
    lines += [("", ""), ("功能团字母", "")]
    for k, v in labels.get("guild_letters", {}).items():
        lines.append((k, v))
    lines += [
        ("", ""),
        ("判定含义", ""),
        ("一致", "源表科/目/纲与 Nemaplex 一致"),
        ("不一致，已按 Nemaplex 修正", "源表与 Nemaplex 不符，输出中采用 Nemaplex 值"),
        ("部分一致", "源表部分字段一致、部分不一致，备注里列出差异"),
        ("Nemaplex 未收录", "该属未出现在 Nemaplex 属索引中"),
        ("名称未解析，未比对", "第 1 步未能确定拉丁名，无法比对"),
        ("源表未给拉丁科/目名，无法比对", "源表只给了中文科名/目名，未做逐字比对"),
    ]
    for a, b in lines:
        ws4.append([a, b])
    ws4.column_dimensions["A"].width = 34
    ws4.column_dimensions["B"].width = 96
    for row in ws4.iter_rows(min_col=2, max_col=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)
    print(f"  -> {path}")


# --------------------------------------------------------------------------- #
# html
# --------------------------------------------------------------------------- #

CSS = """
:root{--bg:#ffffff;--fg:#1d2328;--muted:#5d6b76;--line:#e2e7ea;--accent:#2f6f4e;
--warn:#fff3cd;--bad:#f8d7da;--ok:#e7f3ec;--panel:#f7f9fa;}
*{box-sizing:border-box}
body{margin:0;padding:32px 28px 64px;background:var(--bg);color:var(--fg);
font:14px/1.7 -apple-system,"Segoe UI","Microsoft YaHei",sans-serif;}
.wrap{max-width:1200px;margin:0 auto}
h1{font-size:23px;margin:0 0 6px}
h2{font-size:16px;margin:32px 0 12px;padding-bottom:6px;border-bottom:1px solid var(--line)}
.sub{color:var(--muted);font-size:13px;margin-bottom:22px}
.cards{display:flex;flex-wrap:wrap;gap:12px;margin:18px 0 6px}
.card{flex:1 1 150px;background:var(--panel);border:1px solid var(--line);border-radius:8px;padding:12px 14px}
.card b{display:block;font-size:22px;color:var(--accent);line-height:1.3}
.card span{font-size:12px;color:var(--muted)}
table{border-collapse:collapse;width:100%;font-size:12.5px;margin-top:6px}
th{background:var(--accent);color:#fff;text-align:left;padding:7px 9px;font-weight:600;white-space:nowrap}
td{border-bottom:1px solid var(--line);padding:6px 9px;vertical-align:top}
tr:nth-child(even) td{background:#fbfcfc}
td.warn{background:var(--warn)}
td.bad{background:var(--bad)}
.tag{display:inline-block;padding:1px 7px;border-radius:10px;font-size:11px;background:var(--ok);color:#22593c}
.tag.w{background:var(--warn);color:#7a5b00}
.tag.b{background:var(--bad);color:#842029}
.note{color:var(--muted);font-size:12.5px}
dl{margin:0}
dt{font-weight:600;margin-top:10px}
dd{margin:2px 0 0;color:var(--muted)}
a{color:#0563c1}
"""


def esc(s) -> str:
    return _html.escape(str(s if s is not None else ""))


def write_html(rows: list[dict], labels: dict, meta: dict, stats: dict, path: str) -> None:
    flagged = [r for r in rows if r.get("confidence") in NEEDS_HUMAN]
    guilds = collections.Counter(r["functional_guild"] for r in rows if r.get("functional_guild"))
    cp_ref = {r["cp"]: r for r in rows if r.get("cp")}
    cps = collections.Counter(r["cp"] for r in rows if r.get("cp"))
    fg_ref = {r["feeding_group"]: r for r in rows if r.get("feeding_group")}
    fgs = collections.Counter(r["feeding_group"] for r in rows if r.get("feeding_group"))

    p = []
    p.append('<!DOCTYPE html><html lang="zh-CN"><head><meta charset="utf-8">')
    p.append('<meta name="viewport" content="width=device-width,initial-scale=1">')
    p.append("<title>线虫分类核查报告</title><style>" + CSS + "</style></head><body><div class='wrap'>")
    p.append("<h1>线虫分类核查报告</h1>")
    p.append(f"<div class='sub'>数据源 Nemaplex, UC Davis · 快照 {esc(meta.get('built_at',''))} · "
             f"属索引 {esc(meta.get('genus_count',''))} 属 · 生成于 "
             f"{_dt.datetime.now().strftime('%Y-%m-%d %H:%M')}</div>")

    p.append("<div class='cards'>")
    for label, val in [
        ("处理行数", stats.get("rows", len(rows))),
        ("有分类归属", stats.get("with_classification", 0)),
        ("有 c-p 值", stats.get("with_cp", 0)),
        ("有功能团", stats.get("with_guild", 0)),
        ("待人工确认", len(flagged)),
    ]:
        p.append(f"<div class='card'><b>{esc(val)}</b><span>{esc(label)}</span></div>")
    p.append("</div>")

    p.append("<h2>判定分布</h2><table><tr><th>判定</th><th>行数</th></tr>")
    for k, v in sorted(stats.get("verdict", {}).items(), key=lambda x: -x[1]):
        p.append(f"<tr><td>{esc(k)}</td><td>{esc(v)}</td></tr>")
    p.append("</table>")

    if guilds:
        p.append("<h2>功能团分布</h2><table><tr><th>功能团</th><th>含义</th><th>属数</th></tr>")
        for g, n in sorted(guilds.items()):
            p.append(f"<tr><td>{esc(g)}</td><td>{esc(gw_cn(g, rows))}</td><td>{esc(n)}</td></tr>")
        p.append("</table>")
    if cps:
        p.append("<h2>c-p 值分布</h2><table><tr><th>c-p</th><th>类型</th><th>属数</th></tr>")
        for cp, n in sorted(cps.items()):
            p.append(f"<tr><td>{esc(cp)}</td><td>{esc(cp_ref[cp].get('cp_cn',''))}</td><td>{esc(n)}</td></tr>")
        p.append("</table>")
    if fgs:
        p.append("<h2>取食类群分布</h2><table><tr><th>码</th><th>取食类群</th><th>属数</th></tr>")
        for fg, n in sorted(fgs.items()):
            p.append(f"<tr><td>{esc(fg)}</td><td>{esc(fg_ref[fg].get('feeding_group_cn',''))}</td><td>{esc(n)}</td></tr>")
        p.append("</table>")

    if flagged:
        p.append("<h2>待人工确认</h2><table>")
        p.append("<tr><th>序号</th><th>中文名</th><th>拉丁名</th><th>采用</th><th>置信度</th><th>说明</th></tr>")
        for r in flagged:
            cls = "bad" if r.get("confidence") == "cn-conflict" else "warn"
            p.append(f"<tr><td>{esc(r.get('no'))}</td><td>{esc(r.get('cn_in'))}</td>"
                     f"<td>{esc(r.get('latin_in'))}</td><td>{esc(r.get('latin'))}</td>"
                     f"<td class='{cls}'>{esc(CONF_CN.get(r.get('confidence'), r.get('confidence')))}</td>"
                     f"<td>{esc(r.get('note'))}</td></tr>")
        p.append("</table>")

    p.append("<h2>全部结果</h2><table>")
    p.append("<tr><th>序号</th><th>中文名</th><th>采用拉丁名</th><th>科</th><th>目</th><th>纲</th>"
             "<th>c-p</th><th>取食类群</th><th>功能团</th><th>判定</th></tr>")
    for r in rows:
        cls = ""
        if r.get("confidence") == "cn-conflict":
            cls = " class='bad'"
        elif r.get("confidence") in NEEDS_HUMAN or r.get("verdict") == "not-found":
            cls = " class='warn'"
        link = r.get("genus_url") or ""
        latin = esc(r.get("latin"))
        if link:
            latin = f"<a href='{esc(link)}' target='_blank' rel='noopener'>{latin}</a>"
        p.append(f"<tr{cls}><td>{esc(r.get('no'))}</td><td>{esc(r.get('cn_in'))}</td><td>{latin}</td>"
                 f"<td>{esc(r.get('fam_nem'))}</td><td>{esc(r.get('ord_nem') or r.get('ord_cn'))}</td>"
                 f"<td>{esc(r.get('cls_nem'))}</td><td>{esc(r.get('cp'))}</td>"
                 f"<td>{esc(r.get('feeding_group_cn') or r.get('feeding_group'))}</td>"
                 f"<td>{esc(r.get('functional_guild'))}</td><td>{esc(r.get('verdict_cn'))}</td></tr>")
    p.append("</table>")

    p.append("<h2>说明</h2><dl>")
    p.append("<dt>数据出处</dt><dd>分类与功能团字段全部来自 Nemaplex（http://nemaplex.ucdavis.edu/）的离线快照；"
             "中文译名来自本地对照表，非站点内容。</dd>")
    p.append("<dt>现代体系</dt><dd>De Ley &amp; Blaxter（SSU rDNA）：纲 = Chromadorea 色矛纲 / Enoplea 刺嘴纲；"
             "Tylenchida 垫刃目并入 Rhabditida 小杆目。</dd>")
    p.append("<dt>经典体系</dt><dd>Chitwood (1958)：Adenophorea 无尾感器纲 / Secernentea 尾感器纲。"
             "输出同时给出两套，便于与旧文献对齐。</dd>")
    p.append("<dt>功能团</dt><dd>取食类群首字母 + c-p 值，例如 <code>b1</code> 表示细菌食、极机会型。</dd>")
    p.append("<dt>待人工确认</dt><dd>字典未收录的中文名、疑似拼写错误、以及与源表冲突的行都会标出，"
             "不会静默采用。</dd>")
    p.append("</dl></div></body></html>")

    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(p))
    print(f"  -> {path}")


def gw_cn(guild: str, rows: list[dict]) -> str:
    for r in rows:
        if r.get("functional_guild") == guild:
            return r.get("functional_guild_cn", "")
    return ""


# --------------------------------------------------------------------------- #

def main() -> int:
    ap = argparse.ArgumentParser(description="Render the deliverables (step 4).")
    ap.add_argument("--input", required=True)
    ap.add_argument("--xlsx", default=os.path.join(ROOT, "out", "线虫分类核查结果.xlsx"))
    ap.add_argument("--html", default=os.path.join(ROOT, "out", "report.html"))
    ap.add_argument("--no-html", action="store_true")
    args = ap.parse_args()

    with open(args.input, encoding="utf-8") as fh:
        rows = json.load(fh)

    out_dir = os.path.dirname(args.input)
    stats_path = os.path.join(out_dir, "stats.json")
    stats = load_json(stats_path) if os.path.exists(stats_path) else {}
    labels = load_json(os.path.join(DATA, "zh_cn_labels.json"))
    meta_path = os.path.join(DATA, "_meta.json")
    meta = load_json(meta_path) if os.path.exists(meta_path) else {}

    xlsx_written = False
    if importlib.util.find_spec("openpyxl") is None:
        print("!! 缺少依赖 openpyxl，已跳过 Excel 输出（HTML 报告不受影响）。\n"
              "   安装后重跑同一命令即可得到 Excel 版核查表：\n"
              "   pip install openpyxl", file=sys.stderr)
    else:
        write_xlsx(rows, labels, meta, args.xlsx)
        xlsx_written = True
    if not args.no_html:
        write_html(rows, labels, meta, stats, args.html)
    if not xlsx_written and args.no_html:
        print("!! 注意：--no-html 且未安装 openpyxl，本次运行没有产生任何输出文件。",
              file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
